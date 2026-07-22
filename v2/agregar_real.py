# -*- coding: utf-8 -*-
"""
agregar_real.py — Agrega nueva info REAL al dashboard, en un solo paso.
=======================================================================
Toma un Excel de uploads/real_total/, lo AJUSTA monetariamente a mes 2027 (USD)
y lo AGREGA al maestro de Reales que corresponda (Corp o Distribuibles), para que
el dashboard lo lea en la próxima reconstrucción.

Maneja los dos layouts vistos:
  · Hoja «Consolidado» con columna «Año».
  · Una hoja por año (el año sale del nombre de la hoja; sin columna «Año»).

Rutea por compañía (prefijo de CECO): 1001/1002/1003/1005 → Distribuibles; resto → Corp.
Evita duplicados: si un CECO ya está en el maestro, omite sus filas (re-correr es seguro).

Uso:  python v2/agregar_real.py "uploads/real_total/Mi Real.xlsx"
"""
import os
import sys
import warnings
from collections import Counter

import pandas as pd
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from valor_mes_2027 import cargar_vector, VECTOR_XLSX, VECTOR_SHEET  # noqa: E402

warnings.filterwarnings("ignore")
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

HERE = os.path.dirname(os.path.abspath(__file__))
UP = os.path.join(HERE, "uploads")

# ─────────────────────────── CONFIGURACIÓN ───────────────────────────
CORP_MASTER = os.path.join(UP, "Reales Históricos v2.xlsx")
DIST_MASTER = os.path.join(UP, "Reales Distribuibles Histórico v2.xlsx")
MASTER_SHEET = "Data Consolidada"
DIST_PREFIXES = {"1001", "1002", "1003", "1005"}   # compañías distribuibles
ANIO_DESTINO = 2027

# Columnas del ORIGEN (por nombre)
COL_MES  = "Período"
COL_CECO = "Centro de coste"
COL_VSOC = "Val/Mon.so.CO"
COL_VTR  = "Valor/Mon.tr."
COL_MON  = "Moneda transacción"
COL_ANIO = "Año"

# Mapeo maestro → clave del registro procesado. ('col',k)=copia · ('txt',k)=texto ·
# ('calc',k)=derivado · ('fijo',v)=constante.  El registro procesado usa estas claves:
#   Período, Centro de coste, Val/Mon.so.CO, Texto de pedido, Descrip.clases coste,
#   Denom.clase de coste, Denom.cuenta contrapartida, Moneda sociedad CO, Denominación,
#   Documento compras, Valor/Mon.tr., Moneda transacción, Tipo de Cambio, Unidad TC, Año,
#   Mes gasto, Factor 2027, Valor mes 2027, Unidad 2027, Valor mes 2027 (USD),
#   Valor 2027 USD alt., Factor CPI
_COMUNES = [
    ("periodo_str", "Período"), ("col", "Centro de coste"), ("col", "Val/Mon.so.CO"),
    ("txt", "Texto de pedido"), ("txt", "Descrip.clases coste"), ("txt", "Denom.clase de coste"),
    ("txt", "Denom.cuenta contrapartida"), ("txt", "Moneda sociedad CO"), ("txt", "Denominación"),
    ("txt", "Documento compras"), ("col", "Valor/Mon.tr."), ("txt", "Moneda transacción"),
]
MAP_CORP = _COMUNES + [
    ("calc", "Tipo de Cambio"), ("calc", "Unidad TC"), ("col", "Año"), ("col", "Mes gasto"),
    ("col", "Factor 2027"), ("col", "Valor mes 2027"), ("txt", "Unidad 2027"),
    ("col", "Valor mes 2027 (USD)"), ("col", "Valor 2027 USD alt."), ("col", "Factor CPI"),
    ("fijo", "Real"),
]
MAP_DIST = _COMUNES + [
    ("col", "Año"), ("col", "Mes gasto"), ("col", "Factor 2027"), ("col", "Valor mes 2027"),
    ("txt", "Unidad 2027"), ("col", "Valor mes 2027 (USD)"), ("col", "Valor 2027 USD alt."),
    ("col", "Factor CPI"), ("fijo", "Real"),
]
# ──────────────────────────────────────────────────────────────────────


def _int(x):
    return int(float(x)) if x is not None and pd.notna(x) else None


def calcular(df, anio, vec):
    """Agrega columnas de ajuste 2027 a df. `anio`=None → usa la columna «Año» del df."""
    dolar, ipc_clp, cpi, clpuf = vec["dolar"], vec["ipc_clp"], vec["cpi"], vec["clpuf"]
    get = lambda d, k: (d.get(k) if k is not None else None)

    mes = df[COL_MES].apply(_int)
    yrs = df[COL_ANIO].apply(_int) if anio is None else [anio] * len(df)
    df["Año"] = list(yrs)
    df["Mes gasto"] = [f"{y}.{m:02d}" if (y and m) else None for y, m in zip(df["Año"], mes)]
    df["Mes 2027"] = [f"{ANIO_DESTINO}.{m:02d}" if m else None for m in mes]

    def factor(row):
        mon = str(row[COL_MON]).strip().upper() if pd.notna(row[COL_MON]) else ""
        mg, m27 = row["Mes gasto"], row["Mes 2027"]
        if mon == "USD":
            a, b = get(cpi, m27), get(cpi, mg);       return a / b if a and b else None
        if mon == "CLP":
            a, b = get(ipc_clp, m27), get(ipc_clp, mg); return a / b if a and b else None
        if mon == "UF":
            return get(clpuf, m27)
        return None

    df["Factor 2027"] = df.apply(factor, axis=1)
    df["Valor mes 2027"] = df[COL_VTR] * df["Factor 2027"]

    def unidad(row):
        mon = str(row[COL_MON]).strip().upper() if pd.notna(row[COL_MON]) else ""
        if mon == "UF":
            return "CLP"
        if pd.isna(row["Factor 2027"]):
            return ""
        return mon if mon else ""

    df["Unidad 2027"] = df.apply(unidad, axis=1)
    df["Factor CPI"] = [(get(cpi, m27) / get(cpi, mg)) if get(cpi, m27) and get(cpi, mg) else None
                        for m27, mg in zip(df["Mes 2027"], df["Mes gasto"])]
    df["Valor 2027 USD alt."] = df[COL_VSOC] * df["Factor CPI"]

    def resultado(row):
        u = row["Unidad 2027"]
        if u == "" or pd.isna(u):
            return row["Valor 2027 USD alt."]
        if u == "USD":
            return row["Valor mes 2027"]
        if u == "CLP":
            d = get(dolar, row["Mes 2027"]);  return row["Valor mes 2027"] / d if d else None
        return row["Valor 2027 USD alt."]

    df["Valor mes 2027 (USD)"] = df.apply(resultado, axis=1)
    return df


def leer_y_calcular(path, vec):
    """Lee todas las hojas de datos y devuelve un DataFrame único ya calculado."""
    xls = pd.ExcelFile(path)
    if "Consolidado" in xls.sheet_names:
        hojas = ["Consolidado"]
    else:
        hojas = [s for s in xls.sheet_names if str(s).strip().isdigit() and len(str(s).strip()) == 4]
    if not hojas:
        raise SystemExit("No encontré hoja «Consolidado» ni hojas por año (AAAA) en el archivo.")
    frames = []
    for h in hojas:
        df = pd.read_excel(path, sheet_name=h)
        anio = None if COL_ANIO in df.columns else int(str(h).strip())
        df = calcular(df, anio, vec)
        frames.append(df)
        print(f"    hoja «{h}»: {len(df)} filas (año {'columna' if anio is None else anio})")
    return pd.concat(frames, ignore_index=True)


def _val(kind, key, row):
    if kind == "col":
        v = row.get(key)
        return None if pd.isna(v) else (v.item() if hasattr(v, "item") else v)
    if kind == "txt":
        v = row.get(key)
        return "" if pd.isna(v) else str(v)
    if kind == "periodo_str":
        v = row.get("Período")
        return "" if pd.isna(v) else str(int(float(v)))
    if kind == "fijo":
        return key
    if kind == "calc":
        vsoc, vtr = row.get(COL_VSOC), row.get(COL_VTR)
        mon = str(row.get(COL_MON)).strip().upper() if pd.notna(row.get(COL_MON)) else ""
        if key == "Tipo de Cambio":
            return (vtr / vsoc) if (pd.notna(vtr) and pd.notna(vsoc) and vsoc) else None
        if key == "Unidad TC":
            return f"{mon}/USD" if mon else ""
    return None


def agregar_a_maestro(master_path, mapping, sub):
    """Agrega las filas de `sub` (DataFrame) al maestro, omitiendo CECOs ya presentes."""
    nombre = os.path.basename(master_path)
    print(f"  → {nombre}: {len(sub)} filas candidatas")
    ncols = len(mapping)
    # data_only=True: aplana las fórmulas del maestro a su VALOR en caché. Es CLAVE: sin esto,
    # openpyxl reescribe las fórmulas sin caché y la columna Ajustada («Valor mes 2027 (USD)»)
    # queda vacía al leerla con data_only=True desde construir_v2.
    wb = openpyxl.load_workbook(master_path, data_only=True)
    ws = wb[MASTER_SHEET]
    if ws.max_column != ncols:
        print(f"    OJO: el maestro tiene {ws.max_column} cols; el mapeo espera {ncols}. Se omite.")
        return 0, 0
    existentes = set()
    for (v,) in ws.iter_rows(min_row=2, min_col=2, max_col=2, values_only=True):
        if v is not None:
            existentes.add(str(v))
    nuevos = sub[~sub[COL_CECO].astype(str).isin(existentes)]
    omitidos = len(sub) - len(nuevos)
    cecos_omit = sorted(set(sub[COL_CECO].astype(str)) & existentes)
    if cecos_omit:
        print(f"    omitidos {omitidos} filas de CECOs ya presentes: {cecos_omit}")
    if len(nuevos) == 0:
        print("    nada nuevo para agregar.")
        return 0, omitidos
    for _, row in nuevos.iterrows():
        ws.append([_val(k, key, row) for k, key in mapping])
    wb.save(master_path)
    print(f"    agregadas {len(nuevos)} filas.")
    return len(nuevos), omitidos


def main():
    if len(sys.argv) < 2:
        raise SystemExit("Uso: python v2/agregar_real.py \"<archivo.xlsx>\"")
    path = sys.argv[1]
    if not os.path.isabs(path):
        path = os.path.join(HERE, path) if os.path.exists(os.path.join(HERE, path)) else os.path.abspath(path)
    if not os.path.isfile(path):
        raise SystemExit("No existe el archivo: %s" % path)

    print("Leyendo vector de factores…")
    vec = cargar_vector(VECTOR_XLSX, VECTOR_SHEET)
    print("Procesando (ajuste a 2027):", os.path.basename(path))
    df = leer_y_calcular(path, vec)
    print(f"  Total: {len(df)} filas · sin resultado: {int(df['Valor mes 2027 (USD)'].isna().sum())}")

    # Rutea por compañía (prefijo de CECO)
    pref = df[COL_CECO].astype(str).str[:4]
    es_dist = pref.isin(DIST_PREFIXES)
    corp, dist = df[~es_dist], df[es_dist]
    print(f"  Ruteo: Corp {len(corp)} filas · Distribuibles {len(dist)} filas")
    print("  CECOs:", dict(Counter(df[COL_CECO].astype(str))))

    tot_add = tot_omit = 0
    if len(corp):
        a, o = agregar_a_maestro(CORP_MASTER, MAP_CORP, corp); tot_add += a; tot_omit += o
    if len(dist):
        a, o = agregar_a_maestro(DIST_MASTER, MAP_DIST, dist); tot_add += a; tot_omit += o

    print(f"\n✔ LISTO: {tot_add} filas agregadas al maestro"
          + (f" ({tot_omit} omitidas por CECO ya presente)" if tot_omit else "") + ".")
    print("  Ahora reconstruí/publicá el dashboard para verlo (pestaña «Publicar»).")


if __name__ == "__main__":
    main()
