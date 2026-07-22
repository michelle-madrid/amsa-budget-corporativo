# -*- coding: utf-8 -*-
"""
Calcula, para cada fila de la base de gastos, el «Valor mes 2027 (USD)»:
el gasto llevado a su MISMO mes del año 2027 y expresado en dólares.

Genera columnas auxiliares (una por paso) para poder auditar el cálculo y
guarda el resultado en el mismo Excel, conservando el resto de hojas.

Uso:  python v2/valor_mes_2027.py
"""
import warnings
import pandas as pd
import openpyxl

warnings.filterwarnings("ignore")

# ─────────────────────────── CONFIGURACIÓN ───────────────────────────
# Rutas de archivo
BASE_XLSX   = "v2/uploads/real_total/Real_1000AC3510.xlsx"   # base de gastos (entrada = salida)
VECTOR_XLSX = "v2/uploads/Vector Ajuste Data Histórica.xlsx" # vector de factores

# Hojas
BASE_SHEET   = "Consolidado"
VECTOR_SHEET = "Hoja1"

# Columnas de la base (POR NOMBRE de encabezado; las letras reales van entre paréntesis)
COL_ANIO    = "Año"                  # (A) año del gasto
COL_MES     = "Período"             # (B) número de mes (1..12)
COL_VAL_SOC = "Val/Mon.so.CO"       # (D) valor en moneda sociedad, en USD
COL_VAL_TR  = "Valor/Mon.tr."       # (L) valor en la moneda de la transacción
COL_MON_TR  = "Moneda transacción"  # (M) moneda de la transacción: CLP / USD / UF

# Vector de factores: se lee con header=None (por POSICIÓN, como en Excel).
# Índices de fila 1-based tal como se ven en Excel:
VEC_ROW_PERIODOS = 1   # fila 1: períodos con formato AAAA.MM (ej. 2027.03)
VEC_ROW_DOLAR    = 3   # fila 3: Dólar (CLP/USD)
VEC_ROW_IPC_CLP  = 4   # fila 4: IPC Empalmado CLP
VEC_ROW_CPI      = 6   # fila 6: CPI (USD)
VEC_ROW_CLPUF    = 7   # fila 7: CLP/UF
VEC_COL_INICIO   = 3   # columna 1-based donde empiezan los períodos (C)

ANIO_DESTINO = 2027    # año al que se lleva el gasto

# Nombres de las columnas de resultado (una por paso, para auditar)
C_MES_GASTO = "Mes gasto"
C_MES_2027  = "Mes 2027"
C_FACTOR    = "Factor 2027"
C_VALOR     = "Valor mes 2027"
C_UNIDAD    = "Unidad 2027"
C_FACTOR_CPI= "Factor CPI"
C_ALT       = "Valor 2027 USD alt."
C_RESULTADO = "Valor mes 2027 (USD)"
# ──────────────────────────────────────────────────────────────────────


def norm_periodo(p):
    """Normaliza un período a 'AAAA.MM' (2 dígitos de mes). Acepta str o número."""
    if p is None or (isinstance(p, float) and pd.isna(p)):
        return None
    if isinstance(p, str):
        a, m = p.strip().split(".")
        return f"{int(a)}.{int(m):02d}"
    a = int(p)
    m = round((float(p) - a) * 100)
    return f"{a}.{m:02d}"


def cargar_vector(path, sheet):
    """Lee el vector con header=None y arma un dict período→valor por cada fila relevante."""
    v = pd.read_excel(path, sheet_name=sheet, header=None)
    periodos = [norm_periodo(x) for x in v.iloc[VEC_ROW_PERIODOS - 1, VEC_COL_INICIO - 1:]]

    def fila_dict(row_1based):
        vals = v.iloc[row_1based - 1, VEC_COL_INICIO - 1:]
        return {per: val for per, val in zip(periodos, vals) if per is not None and pd.notna(val)}

    return {
        "dolar":   fila_dict(VEC_ROW_DOLAR),
        "ipc_clp": fila_dict(VEC_ROW_IPC_CLP),
        "cpi":     fila_dict(VEC_ROW_CPI),
        "clpuf":   fila_dict(VEC_ROW_CLPUF),
    }


def main():
    print("Leyendo base de gastos…", BASE_XLSX, "/", BASE_SHEET)
    df = pd.read_excel(BASE_XLSX, sheet_name=BASE_SHEET)
    print(f"  {len(df)} filas · columnas: {list(df.columns)}")

    print("Leyendo vector de factores…", VECTOR_XLSX, "/", VECTOR_SHEET)
    vec = cargar_vector(VECTOR_XLSX, VECTOR_SHEET)
    dolar, ipc_clp, cpi, clpuf = vec["dolar"], vec["ipc_clp"], vec["cpi"], vec["clpuf"]
    print(f"  períodos disponibles: {len(cpi)} (CPI) · ej. rango "
          f"{min(cpi)}..{max(cpi)}")

    # Paso 1 y 2 — claves de período (mes del gasto y su mismo mes en 2027)
    mes = df[COL_MES].apply(lambda x: int(x) if pd.notna(x) else None)
    anio = df[COL_ANIO].apply(lambda x: int(x) if pd.notna(x) else None)
    df[C_MES_GASTO] = [f"{a}.{m:02d}" if a is not None and m is not None else None
                       for a, m in zip(anio, mes)]
    df[C_MES_2027] = [f"{ANIO_DESTINO}.{m:02d}" if m is not None else None for m in mes]

    def get(d, k):
        return d.get(k) if k is not None else None

    # Paso 3 — Factor 2027 según la moneda de transacción
    def factor_2027(row):
        mon = str(row[COL_MON_TR]).strip().upper() if pd.notna(row[COL_MON_TR]) else ""
        mg, m27 = row[C_MES_GASTO], row[C_MES_2027]
        if mon == "USD":
            a, b = get(cpi, m27), get(cpi, mg)
            return a / b if a and b else None
        if mon == "CLP":
            a, b = get(ipc_clp, m27), get(ipc_clp, mg)
            return a / b if a and b else None
        if mon == "UF":
            return get(clpuf, m27)          # la UF ya trae la inflación → se pasa a CLP
        return None                          # otra moneda → vacío

    df[C_FACTOR] = df.apply(factor_2027, axis=1)

    # Paso 4 — Valor mes 2027 = Valor/Mon.tr. × Factor 2027
    df[C_VALOR] = df[COL_VAL_TR] * df[C_FACTOR]

    # Paso 5 — Unidad 2027
    def unidad_2027(row):
        mon = str(row[COL_MON_TR]).strip().upper() if pd.notna(row[COL_MON_TR]) else ""
        if mon == "UF":
            return "CLP"
        if pd.isna(row[C_FACTOR]):
            return ""                        # sin factor → vacío
        return mon if mon else ""

    df[C_UNIDAD] = df.apply(unidad_2027, axis=1)

    # Paso 6 — Factor CPI (aplica a TODAS las filas; la moneda sociedad está en USD)
    df[C_FACTOR_CPI] = [
        (get(cpi, m27) / get(cpi, mg)) if get(cpi, m27) and get(cpi, mg) else None
        for m27, mg in zip(df[C_MES_2027], df[C_MES_GASTO])
    ]

    # Paso 7 — Valor 2027 USD alternativo = Val/Mon.so.CO × Factor CPI
    df[C_ALT] = df[COL_VAL_SOC] * df[C_FACTOR_CPI]

    # Paso 8 — Resultado final
    def resultado(row):
        u = row[C_UNIDAD]
        if u == "" or pd.isna(u):
            return row[C_ALT]                         # sin unidad → alternativo (paso 7)
        if u == "USD":
            return row[C_VALOR]                       # ya en USD
        if u == "CLP":
            d = get(dolar, row[C_MES_2027])           # CLP → USD con dólar de 2027
            return row[C_VALOR] / d if d else None
        return row[C_ALT]

    df[C_RESULTADO] = df.apply(resultado, axis=1)

    # ── Guardar en el MISMO Excel, conservando las demás hojas ──
    aux_cols = [C_MES_GASTO, C_MES_2027, C_FACTOR, C_VALOR, C_UNIDAD,
                C_FACTOR_CPI, C_ALT, C_RESULTADO]
    print("Escribiendo columnas auxiliares en", BASE_SHEET, "…")
    wb = openpyxl.load_workbook(BASE_XLSX)
    ws = wb[BASE_SHEET]
    start = ws.max_column + 1
    for j, name in enumerate(aux_cols):
        ws.cell(row=1, column=start + j, value=name)
        for i, val in enumerate(df[name].tolist()):
            if pd.isna(val):
                continue
            ws.cell(row=i + 2, column=start + j,
                    value=(val.item() if hasattr(val, "item") else val))
    wb.save(BASE_XLSX)
    print(f"  Guardado: {BASE_XLSX} (hoja {BASE_SHEET}, {len(aux_cols)} columnas nuevas)")

    # ── Verificación ──
    verif = [COL_ANIO, COL_MES, COL_VAL_SOC, COL_VAL_TR, COL_MON_TR] + aux_cols
    pd.set_option("display.max_columns", None, "display.width", 240,
                  "display.float_format", lambda x: f"{x:,.2f}")
    print("\n=== head() de verificación ===")
    print(df[verif].head(12).to_string(index=False))

    # Chequeo por moneda (una fila de ejemplo de cada tipo)
    print("\n=== ejemplo por moneda de transacción ===")
    for mon in ["USD", "CLP", "UF"]:
        sub = df[df[COL_MON_TR].astype(str).str.upper() == mon]
        if len(sub):
            print(sub[verif].head(1).to_string(index=False))


if __name__ == "__main__":
    main()
