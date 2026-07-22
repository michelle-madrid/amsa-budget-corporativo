# -*- coding: utf-8 -*-
"""
Ajuste monetario a mes 2027 para archivos de Reales con UNA HOJA POR AÑO
(el año se toma del NOMBRE de la hoja; no hay columna «Año»).

Calcula las mismas columnas auxiliares que valor_mes_2027.py y las escribe de
vuelta en cada hoja de año. Reutiliza la carga del vector de factores.

Uso:  python v2/procesar_real_anios.py
"""
import os
import sys
import warnings

import pandas as pd
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from valor_mes_2027 import (norm_periodo, cargar_vector, VECTOR_XLSX, VECTOR_SHEET,
                            C_MES_GASTO, C_MES_2027, C_FACTOR, C_VALOR, C_UNIDAD,
                            C_FACTOR_CPI, C_ALT, C_RESULTADO)

warnings.filterwarnings("ignore")

# ─────────────────────────── CONFIGURACIÓN ───────────────────────────
FILES = [
    "v2/uploads/real_total/Oficina Londres Real.xlsx",
    "v2/uploads/real_total/Exploraciones Real.xlsx",
]

# Columnas de la base (POR NOMBRE). Estos archivos NO traen «Año» (viene de la hoja).
COL_MES     = "Período"             # número de mes (1..12)
COL_VAL_SOC = "Val/Mon.so.CO"       # valor en moneda sociedad, en USD
COL_VAL_TR  = "Valor/Mon.tr."       # valor en la moneda de la transacción
COL_MON_TR  = "Moneda transacción"  # CLP / USD / UF / (otra)

ANIO_DESTINO = 2027
AUX_COLS = [C_MES_GASTO, C_MES_2027, C_FACTOR, C_VALOR, C_UNIDAD,
            C_FACTOR_CPI, C_ALT, C_RESULTADO]
# ──────────────────────────────────────────────────────────────────────


def _int(x):
    return int(float(x)) if x is not None and pd.notna(x) else None


def calcular(df, anio, vec):
    """Agrega las columnas auxiliares a df (año escalar). Devuelve df."""
    dolar, ipc_clp, cpi, clpuf = vec["dolar"], vec["ipc_clp"], vec["cpi"], vec["clpuf"]
    get = lambda d, k: (d.get(k) if k is not None else None)

    mes = df[COL_MES].apply(_int)
    df[C_MES_GASTO] = [f"{anio}.{m:02d}" if m is not None else None for m in mes]
    df[C_MES_2027]  = [f"{ANIO_DESTINO}.{m:02d}" if m is not None else None for m in mes]

    def factor_2027(row):
        mon = str(row[COL_MON_TR]).strip().upper() if pd.notna(row[COL_MON_TR]) else ""
        mg, m27 = row[C_MES_GASTO], row[C_MES_2027]
        if mon == "USD":
            a, b = get(cpi, m27), get(cpi, mg);      return a / b if a and b else None
        if mon == "CLP":
            a, b = get(ipc_clp, m27), get(ipc_clp, mg); return a / b if a and b else None
        if mon == "UF":
            return get(clpuf, m27)          # la UF ya trae la inflación → pasa a CLP
        return None                          # otra moneda → vacío

    df[C_FACTOR] = df.apply(factor_2027, axis=1)
    df[C_VALOR]  = df[COL_VAL_TR] * df[C_FACTOR]

    def unidad_2027(row):
        mon = str(row[COL_MON_TR]).strip().upper() if pd.notna(row[COL_MON_TR]) else ""
        if mon == "UF":
            return "CLP"
        if pd.isna(row[C_FACTOR]):
            return ""
        return mon if mon else ""

    df[C_UNIDAD] = df.apply(unidad_2027, axis=1)
    df[C_FACTOR_CPI] = [
        (get(cpi, m27) / get(cpi, mg)) if get(cpi, m27) and get(cpi, mg) else None
        for m27, mg in zip(df[C_MES_2027], df[C_MES_GASTO])
    ]
    df[C_ALT] = df[COL_VAL_SOC] * df[C_FACTOR_CPI]

    def resultado(row):
        u = row[C_UNIDAD]
        if u == "" or pd.isna(u):
            return row[C_ALT]
        if u == "USD":
            return row[C_VALOR]
        if u == "CLP":
            d = get(dolar, row[C_MES_2027]);  return row[C_VALOR] / d if d else None
        return row[C_ALT]

    df[C_RESULTADO] = df.apply(resultado, axis=1)
    return df


def anio_de_hoja(nombre):
    s = str(nombre).strip()
    return int(s) if s.isdigit() and len(s) == 4 else None


def main():
    print("Leyendo vector de factores…", VECTOR_XLSX, "/", VECTOR_SHEET)
    vec = cargar_vector(VECTOR_XLSX, VECTOR_SHEET)
    print(f"  períodos CPI: {len(vec['cpi'])}")

    for path in FILES:
        print("\n========================", path)
        xls = pd.ExcelFile(path)
        hojas_anio = [h for h in xls.sheet_names if anio_de_hoja(h) is not None]
        print("  hojas de año:", hojas_anio)

        # 1) calcular por hoja
        calculadas = {}
        for h in hojas_anio:
            df = pd.read_excel(path, sheet_name=h)
            df = calcular(df, anio_de_hoja(h), vec)
            calculadas[h] = df
            print(f"    {h}: {len(df)} filas · sin resultado: {df[C_RESULTADO].isna().sum()}")

        # 2) escribir las columnas auxiliares de vuelta en cada hoja (sin duplicar si ya existen)
        wb = openpyxl.load_workbook(path)
        for h, df in calculadas.items():
            ws = wb[h]
            hdr = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}
            for name in AUX_COLS:
                col = hdr.get(name, ws.max_column + 1)
                if name not in hdr:
                    ws.cell(row=1, column=col, value=name)
                    hdr[name] = col
                for i, val in enumerate(df[name].tolist()):
                    ws.cell(row=i + 2, column=col,
                            value=None if pd.isna(val) else (val.item() if hasattr(val, "item") else val))
        wb.save(path)
        print(f"  Guardado (columnas 2027 en {len(calculadas)} hojas): {path}")

        # 3) head de verificación (primer año)
        h0 = hojas_anio[0]
        verif = [COL_MES, COL_VAL_SOC, COL_VAL_TR, COL_MON_TR] + AUX_COLS
        pd.set_option("display.max_columns", None, "display.width", 240,
                      "display.float_format", lambda x: f"{x:,.2f}")
        print(f"  --- head {h0} ---")
        print(calculadas[h0][verif].head(6).to_string(index=False))


if __name__ == "__main__":
    main()
