# -*- coding: utf-8 -*-
"""
aplanar_maestro.py — Convierte las FÓRMULAS de un maestro de Reales a sus VALORES.
==================================================================================
Previene el bug de la columna Ajustada en 0: si el maestro tiene fórmulas (con
caché), un openpyxl load+save las rompe. Aplanarlas a valores lo deja a prueba de
futuros agregados.

Hace: respaldo (.BAK) → carga data_only=True (aplana) → guarda → verifica que no
se perdió nada (mismas filas, Ajustada sin nulos, misma suma).

Uso:  python v2/aplanar_maestro.py "uploads/Reales Distribuibles Histórico v2.xlsx"
"""
import os
import sys
import shutil
import warnings

import openpyxl

warnings.filterwarnings("ignore")
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

SHEET = "Data Consolidada"
COL_AJUSTADA = "Valor mes 2027 (USD)"


def stats(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[SHEET]
    it = ws.iter_rows(values_only=True)
    hdr = list(next(it))
    ci = hdr.index(COL_AJUSTADA)
    n = nul = 0
    s = 0.0
    for r in it:
        n += 1
        v = r[ci] if len(r) > ci else None
        if v is None:
            nul += 1
        elif isinstance(v, (int, float)):
            s += v
    wb.close()
    return n, nul, s, ci


def main():
    if len(sys.argv) < 2:
        raise SystemExit("Uso: python v2/aplanar_maestro.py \"<maestro.xlsx>\"")
    path = sys.argv[1]
    if not os.path.isabs(path):
        cand = os.path.join(os.path.dirname(os.path.abspath(__file__)), path)
        path = cand if os.path.exists(cand) else os.path.abspath(path)
    if not os.path.isfile(path):
        raise SystemExit("No existe: %s" % path)

    n0, nul0, s0, ci = stats(path)
    print("ANTES:  %d filas · Ajustada(col %d) nulos=%d suma=%.1fMM" % (n0, ci, nul0, s0 / 1e6))

    bak = path[:-5] + ".BAK.xlsx" if path.lower().endswith(".xlsx") else path + ".BAK"
    if os.path.exists(bak):
        print("Respaldo ya existía:", os.path.basename(bak))
    else:
        shutil.copy2(path, bak)
        print("Respaldo creado:", os.path.basename(bak))

    print("Aplanando (data_only=True + save)…")
    wb = openpyxl.load_workbook(path, data_only=True)
    wb.save(path)
    wb.close()

    n1, nul1, s1, _ = stats(path)
    print("DESPUÉS: %d filas · Ajustada nulos=%d suma=%.1fMM" % (n1, nul1, s1 / 1e6))

    if n1 == n0 and nul1 == 0 and abs(s1 - s0) < 1.0:
        print("✔ OK: aplanado a valores sin pérdida de datos.")
    else:
        print("✖ ATENCIÓN: los números no coinciden. Restaurá desde", os.path.basename(bak))
        sys.exit(1)


if __name__ == "__main__":
    main()
