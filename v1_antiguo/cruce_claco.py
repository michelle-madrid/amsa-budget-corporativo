# -*- coding: utf-8 -*-
"""
cruce_claco.py — Cruce del Diccionario de CLACOS
================================================

Objetivo (reutilizable para futuras cargas de CLACOS.xlsx):

  1) Leer la hoja "Diccionario" de uploads/CLACOS.xlsx.
  2) Cruzar el "Nombre SAP" (nombre truncado con el que vienen los exports de SAP)
     contra su "código" (Cód_Agrupación2, ej. REMUNERACI.EE) y su "Nombre oficial".
  3) Escribir el listado (Código + Nombre oficial [+ Nombre SAP]) en una pestaña
     llamada "Diccionario CLACO" dentro del mismo libro.

Uso:
    python cruce_claco.py                      # usa uploads/CLACOS.xlsx
    python cruce_claco.py otra_ruta.xlsx       # usa otro archivo

No borra ni altera las demás hojas del libro: solo (re)genera "Diccionario CLACO".
"""

import os
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# ---------------------------------------------------------------------------
# Configuración
# ---------------------------------------------------------------------------
BASE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_FILE = os.path.join(BASE, "uploads", "CLACOS.xlsx")

HOJA_ORIGEN = "Diccionario"
HOJA_DESTINO = "Diccionario CLACO"

# Nombres de columna esperados en la hoja Diccionario (se buscan por texto en la
# fila de encabezado, así el script no depende de posiciones fijas).
COL_CODIGO = "Cód_Agrupación2"
COL_NOMBRE = "Nombre oficial"
COL_SAP = "Nombre SAP"


def _norm(v):
    """Normaliza una celda de encabezado a texto comparable."""
    return str(v).strip().lower() if v is not None else ""


def localizar_columnas(ws):
    """Devuelve {clave: índice de columna (0-based)} buscando por encabezado."""
    encabezado = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    objetivo = {
        "codigo": _norm(COL_CODIGO),
        "nombre": _norm(COL_NOMBRE),
        "sap": _norm(COL_SAP),
    }
    idx = {}
    for j, celda in enumerate(encabezado):
        n = _norm(celda)
        for clave, texto in objetivo.items():
            if n == texto:
                idx[clave] = j
    faltan = [k for k in objetivo if k not in idx]
    if faltan:
        raise SystemExit(
            f"No encontré las columnas {faltan} en la hoja '{HOJA_ORIGEN}'. "
            f"Encabezado leído: {encabezado}"
        )
    return idx


def extraer_cruce(ws, idx):
    """
    Recorre las filas de datos y arma el listado cruzado.
    Cada registro: (codigo, nombre_oficial, nombre_sap).
    Se descartan filas sin código NI nombre oficial (filas vacías).
    """
    filas = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        codigo = row[idx["codigo"]] if idx["codigo"] < len(row) else None
        nombre = row[idx["nombre"]] if idx["nombre"] < len(row) else None
        sap = row[idx["sap"]] if idx["sap"] < len(row) else None

        # Limpieza: strings sin espacios extremos; "" -> None
        def limpio(v):
            if v is None:
                return None
            if isinstance(v, str):
                v = v.strip()
                return v or None
            return v

        codigo, nombre, sap = limpio(codigo), limpio(nombre), limpio(sap)

        if codigo is None and nombre is None:
            continue  # fila vacía
        filas.append((codigo, nombre, sap))
    return filas


def escribir_diccionario_claco(wb, filas):
    """(Re)crea la pestaña 'Diccionario CLACO' con Código · Nombre oficial · Nombre SAP."""
    if HOJA_DESTINO in wb.sheetnames:
        del wb[HOJA_DESTINO]
    ws = wb.create_sheet(HOJA_DESTINO)

    encabezados = ["Código", "Nombre oficial", "Nombre SAP"]
    ws.append(encabezados)

    # Estilo de encabezado
    fill = PatternFill("solid", fgColor="1F3864")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(vertical="center")

    for codigo, nombre, sap in filas:
        ws.append([codigo, nombre, sap])

    # Ancho de columnas
    anchos = [20, 45, 24]
    for i, w in enumerate(anchos, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return ws


def main():
    ruta = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_FILE
    if not os.path.exists(ruta):
        raise SystemExit(f"No existe el archivo: {ruta}")

    print(f"Leyendo: {ruta}")
    wb = openpyxl.load_workbook(ruta)  # conserva las demás hojas tal cual
    if HOJA_ORIGEN not in wb.sheetnames:
        raise SystemExit(f"El libro no tiene la hoja '{HOJA_ORIGEN}'. Hojas: {wb.sheetnames}")

    ws = wb[HOJA_ORIGEN]
    idx = localizar_columnas(ws)
    filas = extraer_cruce(ws, idx)

    con_sap = sum(1 for _, _, s in filas if s)
    print(f"  Registros en Diccionario: {len(filas)}")
    print(f"  Con Nombre SAP (cruzados): {con_sap}")
    print(f"  Sin Nombre SAP:           {len(filas) - con_sap}")

    escribir_diccionario_claco(wb, filas)
    wb.save(ruta)
    print(f"OK -> pestaña '{HOJA_DESTINO}' escrita en {ruta}")


if __name__ == "__main__":
    main()
