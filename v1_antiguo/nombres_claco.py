# -*- coding: utf-8 -*-
"""
nombres_claco.py — Nombres oficiales de Ítem Relevante desde el diccionario CLACO
=================================================================================

El panel "Ítem Relevante" del dashboard muestra `itemNames` de data.js, un mapa
   clave = Nombre SAP (truncado a 20)  ->  nombre a mostrar
Este módulo re-sincroniza esos nombres con el "Nombre oficial" del diccionario
CLACO (uploads/CLACOS.xlsx, hoja "Diccionario"), cruzando por el Nombre SAP.

Doble uso:

  1) Reutilizable AHORA (aplica al HTML + data.js ya publicados):
        python nombres_claco.py
     Reinyecta el data.js dentro del HTML, deja el respaldo .bak y regenera el ZIP.

  2) En el PIPELINE (construir_bbdd.py): tras leer los mapas conservados del HTML
     se llama a `aplicar_itemnames()` para que cada regeneración adopte el nombre
     oficial CLACO automáticamente. Ver el enganche en construir_bbdd.py.

Los nombres se toman TAL CUAL vienen en el Excel (sin limpiar cortes/espacios),
según lo pedido.
"""

import os
import gzip
import base64
import json
import shutil
import zipfile
import datetime

import openpyxl

# Reutiliza rutas y helpers de inyección del actualizador (importar NO abre la GUI).
import actualizar_dashboard as ad

HERE = os.path.dirname(os.path.abspath(__file__))
CLACOS_FILE = os.path.join(HERE, "uploads", "CLACOS.xlsx")
HTML_PATH = os.path.join(HERE, ad.HTML_NAME)

HOJA_DIC = "Diccionario"
COL_NOMBRE = "nombre oficial"
COL_SAP = "nombre sap"
COL_COD = "cód_agrupación2"

# Correcciones de Item Relevante en distribuibles 2022-2025 ('full'): el Excel dejó el
# NOMBRE del CECO en la columna "Item Relevante" en vez del ítem. El ítem correcto se
# verificó contra el archivo 2026 FY (col B, código AMSA…). clave = texto errado ->
# ítem correcto (debe ser una clave existente de itemNames, truncada a 20 como en SAP).
ITEM_FIX_DIST = {
    "Comunidades RCA": "Liquidación Interna",    # CECO 1001AD4507 (Asuntos Públicos)
    "Corporativo":     "Servicios Corporativ",   # CECO 1005AD0504 (→ Servicios Corporativos)
}


def corregir_item_dist(item):
    """Devuelve el ítem correcto para un valor de distribuibles (o el mismo)."""
    return ITEM_FIX_DIST.get(item, item)


def _norm(v):
    return str(v).strip().lower() if v is not None else ""


def mapa_oficial(path=CLACOS_FILE):
    """Devuelve {Nombre SAP (str): Nombre oficial} desde la hoja Diccionario.

    Solo incluye filas con Nombre SAP no vacío (son las que cruzan con itemNames).
    Los valores se dejan tal cual (sin normalizar).
    """
    if not os.path.exists(path):
        raise SystemExit(f"No existe el archivo CLACOS: {path}")
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if HOJA_DIC not in wb.sheetnames:
        raise SystemExit(f"El libro no tiene la hoja '{HOJA_DIC}'. Hojas: {wb.sheetnames}")
    ws = wb[HOJA_DIC]
    filas = ws.iter_rows(values_only=True)
    encabezado = next(filas)
    idx = {}
    for j, celda in enumerate(encabezado):
        n = _norm(celda)
        if n == COL_NOMBRE:
            idx["nombre"] = j
        elif n == COL_SAP:
            idx["sap"] = j
    if "nombre" not in idx or "sap" not in idx:
        raise SystemExit(
            f"No encontré 'Nombre oficial' y/o 'Nombre SAP' en '{HOJA_DIC}'. "
            f"Encabezado: {encabezado}"
        )
    mapa = {}
    for row in filas:
        sap = row[idx["sap"]] if idx["sap"] < len(row) else None
        off = row[idx["nombre"]] if idx["nombre"] < len(row) else None
        if isinstance(sap, str):
            sap = sap.strip()
        if not sap:
            continue
        mapa[str(sap)] = off
    wb.close()
    return mapa


def aplicar_itemnames(itemNames, mapa):
    """Sobre-escribe los nombres de `itemNames` cuyo Nombre SAP (la clave) exista en
    el diccionario CLACO. NO agrega claves nuevas: solo re-nombra las ya presentes.

    Devuelve (nuevo_dict, cambios) donde cambios = [(clave, antes, despues), ...].
    """
    nuevo = dict(itemNames)
    cambios = []
    for clave, actual in itemNames.items():
        if clave in mapa and mapa[clave] != actual:
            nuevo[clave] = mapa[clave]
            cambios.append((clave, actual, mapa[clave]))
    return nuevo, cambios


# --------------------------------------------------------------------------
# Códigos (Cód_Agrupación2) por Ítem Relevante
# --------------------------------------------------------------------------
def _mapas_codigo(path=CLACOS_FILE):
    """Devuelve 3 lookups del diccionario CLACO para cruzar ítems → código:
       sap2cod   : Nombre SAP (exacto)          -> Cód_Agrupación2
       off2cod   : Nombre oficial (normalizado) -> Cód_Agrupación2
       trunc2cod : Nombre oficial[:20] (norm)   -> Cód_Agrupación2  (para claves truncadas por SAP)
    """
    if not os.path.exists(path):
        raise SystemExit(f"No existe el archivo CLACOS: {path}")
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[HOJA_DIC]
    it = ws.iter_rows(values_only=True)
    hdr = [_norm(h) for h in next(it)]
    for col in (COL_COD, COL_NOMBRE, COL_SAP):
        if col not in hdr:
            raise SystemExit(f"Falta la columna '{col}' en '{HOJA_DIC}'. Encabezado: {hdr}")
    ci, ni, si = hdr.index(COL_COD), hdr.index(COL_NOMBRE), hdr.index(COL_SAP)
    sap2cod, off2cod, trunc2cod = {}, {}, {}
    for r in it:
        cod = r[ci] if ci < len(r) else None
        off = r[ni] if ni < len(r) else None
        sap = r[si] if si < len(r) else None
        if cod is None:
            continue
        if isinstance(sap, str):
            sap = sap.strip()
        if sap:
            sap2cod[str(sap)] = cod
        if off is not None:
            o = str(off).strip()
            off2cod.setdefault(o.lower(), cod)
            trunc2cod.setdefault(o[:20].strip().lower(), cod)
    wb.close()
    return sap2cod, off2cod, trunc2cod


def construir_itemcodes(item_keys, item_names, path=CLACOS_FILE):
    """Mapa {itemKey: Cód_Agrupación2} para las claves de ítem del dashboard.
    Cruza por Nombre SAP (clave exacta) y, si no, por Nombre oficial (y su truncado
    a 20, para las claves que SAP recorta). Omite las claves que no cruzan.
    """
    sap2cod, off2cod, trunc2cod = _mapas_codigo(path)
    out = {}
    for k in item_keys:
        if k is None:
            continue
        if k in sap2cod:
            out[k] = sap2cod[k]
            continue
        disp = str(item_names.get(k, k)).strip().lower()
        kk = str(k).strip().lower()
        cod = off2cod.get(disp) or trunc2cod.get(kk) or off2cod.get(kk)
        if cod is not None:
            out[k] = cod
    return out


# --------------------------------------------------------------------------
# Aplicación directa al data.js embebido en el HTML (uso standalone)
# --------------------------------------------------------------------------
def _leer_data_js(html):
    m = ad._asset_regex().search(html)
    if not m:
        raise SystemExit("No encontré el bloque de datos (data.js) dentro del HTML.")
    return gzip.decompress(base64.b64decode(m.group(2))).decode("utf-8")


def actualizar_data_js(data_js, mapa, path=CLACOS_FILE):
    """Actualiza los bloques window.CORP_DATA (itemNames + itemCodes) y window.DIST_DATA
    (corrige ítems mal etiquetados vía ITEM_FIX_DIST), conservando intacto todo lo demás
    (DOT_DATA, CORP_DICT, CORP_LOGO…).

    Devuelve (nuevo_data_js, cambios_nombres, itemCodes, n_items_dist_corregidos).
    """
    dec = json.JSONDecoder()

    def loc(var):
        mk = "window." + var + " = "
        i = data_js.find(mk)
        if i == -1:
            return None
        s = i + len(mk)
        obj, e = dec.raw_decode(data_js, s)
        return s, e, obj

    corp = loc("CORP_DATA")
    if corp is None:
        raise SystemExit("No encontré window.CORP_DATA en el data.js.")
    s_c, e_c, ada = corp
    dist = loc("DIST_DATA")   # (s, e, obj) o None

    # 1) Nombres oficiales CLACO en itemNames
    ada["itemNames"], cambios = aplicar_itemnames(ada.get("itemNames", {}), mapa)

    # 2) Corrección de ítems mal etiquetados en distribuibles
    fixes = 0
    if dist is not None:
        for r in dist[2].get("records", []):
            it = r.get("item")
            fix = corregir_item_dist(it)
            if fix != it:
                r["item"] = fix
                fixes += 1

    # 3) itemCodes sobre el universo corregido (CORP items+records + DIST records)
    keys = set(ada.get("items", []))
    for r in ada.get("records", []):
        keys.add(r.get("item"))
    if dist is not None:
        for r in dist[2].get("records", []):
            keys.add(r.get("item"))
    ada["itemCodes"] = construir_itemcodes(keys, ada["itemNames"], path)

    j = lambda o: json.dumps(o, ensure_ascii=False, separators=(",", ":"))
    if dist is not None:
        s_d, e_d, dd = dist
        nuevo = data_js[:s_c] + j(ada) + data_js[e_c:s_d] + j(dd) + data_js[e_d:]
    else:
        nuevo = data_js[:s_c] + j(ada) + data_js[e_c:]
    return nuevo, cambios, ada["itemCodes"], fixes


def main():
    print(f"Leyendo diccionario CLACO: {CLACOS_FILE}")
    mapa = mapa_oficial(CLACOS_FILE)
    print(f"  Nombres oficiales con Nombre SAP: {len(mapa)}")

    print(f"Leyendo HTML: {HTML_PATH}")
    with open(HTML_PATH, "r", encoding="utf-8", newline="") as f:
        html = f.read()
    data_js = _leer_data_js(html)
    nuevo_data_js, cambios, itemcodes, fixes = actualizar_data_js(data_js, mapa)

    print(f"Cambios en itemNames ({len(cambios)}):" if cambios else "itemNames: ya coincidía con los nombres oficiales CLACO.")
    for clave, antes, despues in cambios:
        print(f"  [{clave}]  {antes!r}  ->  {despues!r}")
    print(f"Ítems distribuibles corregidos (registros remapeados): {fixes}")
    print(f"itemCodes generados: {len(itemcodes)} ítems con código.")

    shutil.copyfile(HTML_PATH, HTML_PATH + ".bak")
    print("Respaldo creado: " + os.path.basename(HTML_PATH) + ".bak")
    nuevo_html = ad.inyectar(html, nuevo_data_js)              # solo el asset data.js
    with open(HTML_PATH, "w", encoding="utf-8", newline="") as f:
        f.write(nuevo_html)
    with open(os.path.join(HERE, "data.js"), "w", encoding="utf-8", newline="") as f:
        f.write(nuevo_data_js)
    print("HTML actualizado e inyectado · data.js regenerado.")
    print("NOTA: si editaste app.jsx/model.js/*.jsx, corre además: python reembedir_codigo.py")

    zip_path = os.path.join(HERE, "Dashboard Actividad Corporativa.zip")
    readme = ("DASHBOARD - Actividad Corporativa + Distribuibles (AMSA)\r\n\r\n"
              "Doble clic en 'Dashboard Actividad Corporativa.html' (Chrome/Edge/Firefox).\r\n"
              f"\r\nActualizado: {datetime.date.today().strftime('%d-%m-%Y')}\r\n")
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as z:
        z.write(HTML_PATH, arcname=ad.HTML_NAME)
        z.writestr("LEER - Como abrir.txt", readme)
    print("ZIP listo: " + os.path.basename(zip_path))


if __name__ == "__main__":
    main()
