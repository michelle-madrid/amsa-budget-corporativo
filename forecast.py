# -*- coding: utf-8 -*-
"""
forecast.py — Forecast 5+7 2026 por CECO × CLACO
================================================
Lee `uploads/FORECAST 5+7.XLSX` (hoja Sheet1) y arma el bloque `window.CORP_FCST26`
que consume el dashboard:

  - Agrupa el detalle (7.011 filas) por (CECO, CLACO) sumando el Total-2026.
  - Mapea cada CECO → VP / Gerencia / Resumen (CECOS.xlsx) y cada CLACO → Ítem
    Relevante (CLACO → Cód_Agrupación2 → Nombre oficial del diccionario CLACO).
  - Carga TODOS los CECOs; el campo "Resumen" de CECOS queda como dimensión de
    filtro (Act. Corp. + Distribuibles / Proyectos / Movimiento Financiero / …).

Estructura emitida (compacta; el front hace el join):
  window.CORP_FCST26 = {
    version: "FF05_26V01",
    total: <USD>,
    cecoMeta:  { <ceco>: {cd,comp,res,vp,ger} },     # 1 por CECO
    clacoMeta: { <claco>:{cld,ag,item} },            # 1 por CLACO
    records:   [ {ceco,claco,val} ]                  # solo no-cero
  }

Uso:
  · En el pipeline: construir_bbdd llama a `cargar_forecast(cecos)` y agrega el
    bloque al data.js.
  · Standalone (inyecta en el HTML/data.js vigente sin regenerar todo):
        python forecast.py
"""
import os
import json
import gzip
import base64
import shutil
import zipfile
import datetime
import collections
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
UP = os.path.join(HERE, "uploads")
FCST_FILE = os.path.join(UP, "FORECAST 5+7.XLSX")
CECOS_FILE = os.path.join(UP, "CECOS.xlsx")
CLACOS_FILE = os.path.join(UP, "CLACOS.xlsx")

# Columnas de Sheet1 (0-based)
C_VERSION, C_CECO, C_CECODESC, C_CLACO, C_CLACODESC, C_COMP, C_TOTAL = 0, 2, 3, 4, 6, 9, 23

SIN = "(sin clasificar)"


def _norm(v):
    return str(v).strip().lower() if v is not None else ""


def _resumen_por_ceco(path=CECOS_FILE):
    """{CECO: Resumen} desde la hoja 'CECOS Corporativo' (col Resumen / col CECO)."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb["CECOS Corporativo"]
    rows = ws.iter_rows(values_only=True)
    hdr = [_norm(c) for c in next(rows)]
    ci_res = hdr.index("resumen")
    ci_ceco = hdr.index("ceco")
    out = {}
    for r in rows:
        if r[ci_ceco] is not None:
            out[str(r[ci_ceco]).strip()] = r[ci_res]
    wb.close()
    return out


def _claco_maps(path=CLACOS_FILE):
    """(claco2ag, ag2nom): CLACO→Cód_Agrupación2 (GP) y Cód_Agrupación2→Nombre oficial."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    claco2ag = {}
    for sh, ci_claco, ci_ag in [("GP_CO1", 12, 7), ("GP_CO2P", 6, 2)]:
        rows = wb[sh].iter_rows(values_only=True)
        next(rows)
        for r in rows:
            if r[ci_claco] is not None:
                try:
                    claco2ag.setdefault(int(r[ci_claco]), str(r[ci_ag]).strip())
                except (TypeError, ValueError):
                    pass
    ag2nom = {}
    rows = wb["Diccionario"].iter_rows(values_only=True)
    next(rows)
    for r in rows:
        if r[0] is not None:
            ag2nom[str(r[0]).strip()] = r[1]
    wb.close()
    return claco2ag, ag2nom


def cargar_forecast(cecos, fcst_path=FCST_FILE, cecos_path=CECOS_FILE, clacos_path=CLACOS_FILE):
    """Devuelve el dict CORP_FCST26. `cecos` = {code: (vp, ger, tc, ap)} (de CECOS)."""
    if not os.path.exists(fcst_path):
        raise SystemExit(f"No existe el forecast: {fcst_path}")
    resumen = _resumen_por_ceco(cecos_path)
    claco2ag, ag2nom = _claco_maps(clacos_path)

    wb = openpyxl.load_workbook(fcst_path, data_only=True, read_only=True)
    ws = wb["Sheet1"]
    it = ws.iter_rows(values_only=True)
    next(it)
    version = None
    by_cc = collections.defaultdict(float)                 # (ceco,claco) -> val
    ceco_meta = {}                                          # ceco -> {cd,comp}
    claco_meta = {}                                         # claco -> {cld}
    for r in it:
        if r[C_CECO] is None:
            continue
        version = version or r[C_VERSION]
        ceco = str(r[C_CECO]).strip()
        claco = r[C_CLACO]
        val = r[C_TOTAL] or 0
        by_cc[(ceco, claco)] += val
        if ceco not in ceco_meta:
            ceco_meta[ceco] = {"cd": r[C_CECODESC], "comp": str(r[C_COMP]) if r[C_COMP] is not None else ""}
        if claco not in claco_meta:
            claco_meta[claco] = {"cld": r[C_CLACODESC]}
    wb.close()

    def _r(v):
        f = round(float(v), 2)
        return int(f) if f == int(f) else f

    # Metadata por CECO (VP/Ger/Resumen)
    cecoMeta = {}
    for ceco, m in ceco_meta.items():
        vg = cecos.get(ceco)
        vp = vg[0] if vg else SIN
        ger = vg[1] if vg else (m["cd"] or SIN)
        res = resumen.get(ceco) or SIN
        cecoMeta[ceco] = {"cd": m["cd"] or "", "comp": m["comp"], "res": res, "vp": vp, "ger": ger}

    # Metadata por CLACO (Ítem vía Agrupación2)
    clacoMeta = {}
    for claco, m in claco_meta.items():
        try:
            cl = int(claco)
        except (TypeError, ValueError):
            cl = claco
        ag = claco2ag.get(cl)
        item = (ag2nom.get(ag) if ag else None) or (m["cld"] or SIN)
        clacoMeta[str(claco)] = {"cld": m["cld"] or "", "ag": ag or "", "item": item}

    records = []
    total = 0.0
    for (ceco, claco), val in by_cc.items():
        if round(val, 2) == 0:
            continue
        records.append({"ceco": ceco, "claco": str(claco), "val": _r(val)})
        total += val

    return {
        "version": version,
        "total": _r(total),
        "cecoMeta": cecoMeta,
        "clacoMeta": clacoMeta,
        "records": records,
    }


# --------------------------------------------------------------------------
# Standalone: inyecta window.CORP_FCST26 en el data.js embebido del HTML
# --------------------------------------------------------------------------
def _bloque_js(fcst):
    j = lambda o: json.dumps(o, ensure_ascii=False, separators=(",", ":"))
    return "window.CORP_FCST26 = " + j(fcst) + ";"


def main():
    import actualizar_dashboard as ad
    import construir_bbdd as cb
    html_path = cb.HTML_PATH
    html = open(html_path, "r", encoding="utf-8", newline="").read()
    maps = ad.maps_actuales(html)
    cecos = cb.cecos_por_codigo(maps[1])
    fcst = cargar_forecast(cecos)
    print(f"Forecast {fcst['version']}: {fcst['total']/1e6:,.2f} MM · "
          f"{len(fcst['records'])} registros CECO×CLACO · "
          f"{len(fcst['cecoMeta'])} CECOs · {len(fcst['clacoMeta'])} CLACOs")

    data_js = ad._asset_regex().search(html)
    if not data_js:
        raise SystemExit("No encontré el data.js embebido.")
    dj = gzip.decompress(base64.b64decode(data_js.group(2))).decode("utf-8")

    bloque = _bloque_js(fcst)
    if "window.CORP_FCST26" in dj:
        # reemplaza el bloque existente
        import re
        dj = re.sub(r"window\.CORP_FCST26 = .*?;(?=\nwindow\.|\Z)", bloque, dj, count=1, flags=re.S)
    else:
        # inserta tras el bloque DOT_DATA
        mk = "window.DOT_DATA = "
        i = dj.find(mk)
        if i == -1:
            raise SystemExit("No encontré window.DOT_DATA para insertar el forecast.")
        _, e = json.JSONDecoder().raw_decode(dj, i + len(mk))
        # e apunta al final del objeto; tras él viene ';' y luego el resto (cola)
        semi = dj.find(";", e)
        dj = dj[:semi + 1] + "\n" + bloque + dj[semi + 1:]

    shutil.copyfile(html_path, html_path + ".bak")
    nuevo_html = ad.inyectar(html, dj)
    with open(html_path, "w", encoding="utf-8", newline="") as f:
        f.write(nuevo_html)
    with open(os.path.join(HERE, "data.js"), "w", encoding="utf-8", newline="") as f:
        f.write(dj)
    print("HTML actualizado con window.CORP_FCST26 · data.js regenerado.")
    zip_path = os.path.join(HERE, "Dashboard Actividad Corporativa.zip")
    readme = ("DASHBOARD - Actividad Corporativa + Distribuibles (AMSA)\r\n\r\n"
              "Doble clic en 'Dashboard Actividad Corporativa.html'.\r\n"
              f"\r\nActualizado: {datetime.date.today().strftime('%d-%m-%Y')}\r\n")
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as z:
        z.write(html_path, arcname=ad.HTML_NAME)
        z.writestr("LEER - Como abrir.txt", readme)
    print("ZIP regenerado.")


if __name__ == "__main__":
    main()
