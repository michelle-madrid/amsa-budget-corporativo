# -*- coding: utf-8 -*-
"""
Genera la hoja «… Unpivot» de un ejercicio (Forecast / Outlook / Ppto) dentro del MISMO Excel.

Qué hace
--------
La hoja cruda viene ancha: 14 columnas de identidad + 12 columnas de meses
(Enero-AAAA … Diciembre-AAAA) + Total-AAAA. Este script la pasa a formato largo:

    <14 columnas de identidad> | Mes-Año | Valor | Valor USD <AÑO> (mes equivalente)

  · Mes-Año  → texto "AAAA-MM" (ej. "2026-01")
  · Valor    → el monto del mes (los ceros se conservan; Total-AAAA se ignora)
  · Valor USD <AÑO> (mes equivalente)
               = Valor × CPI(añoDestino-MM) / CPI(añoOrigen-MM)

El CPI sale de la fila «CPI» de «Vector Ajuste Data Histórica.xlsx». La fórmula se
verificó contra «EXPORT_FORECAST_5+7 2026.xlsx» (hoja «Forecast 5+7 Unpivot»): error 0.

Uso
---
    python unpivot_ejercicio.py "uploads/ejercicios_2026/Outlook 6+6 2026.XLSX"
    python unpivot_ejercicio.py "<archivo.xlsx>" --salida "Outlook 6+6 Unpivot"
    python unpivot_ejercicio.py "<archivo.xlsx>" --anio-destino 2028 --hoja Sheet1

Por defecto el nombre de la hoja de salida se deriva del archivo quitándole el año:
«Outlook 6+6 2026.XLSX» → «Outlook 6+6 Unpivot». La hoja se AGREGA al mismo archivo
(si ya existía una con ese nombre, se reemplaza). El resto del libro no se toca.
"""
import argparse
import os
import re
import sys

import openpyxl

# La consola de Windows (cp1252) puede reventar al imprimir «», → o miles: forzamos UTF-8 tolerante.
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

BASE = os.path.dirname(os.path.abspath(__file__))
VECTOR_DEF = os.path.join(BASE, "uploads", "Vector Ajuste Data Histórica.xlsx")

MESES = {
    "enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6,
    "julio": 7, "agosto": 8, "septiembre": 9, "setiembre": 9, "octubre": 10,
    "noviembre": 11, "diciembre": 12,
}


def leer_cpi(path, serie="CPI"):
    """Devuelve {'AAAA.MM': valor} de la fila indicada del vector de ajuste."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    cols = {}
    for c in range(1, ws.max_column + 1):
        h = ws.cell(1, c).value
        if h is not None:
            cols[str(h).strip()] = c
    fila = None
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(r, 1).value or "").strip().upper() == serie.upper():
            fila = r
            break
    if fila is None:
        wb.close()
        raise SystemExit("ERROR: no encontré la fila «%s» en %s" % (serie, os.path.basename(path)))
    out = {k: ws.cell(fila, c).value for k, c in cols.items()}
    wb.close()
    return {k: v for k, v in out.items() if isinstance(v, (int, float))}


def parse_mes(header):
    """'Enero-2026' → (2026, 1). Devuelve None si no es una columna de mes."""
    if not isinstance(header, str):
        return None
    m = re.match(r"^\s*([A-Za-zÁÉÍÓÚáéíóúñÑ]+)\s*-\s*(\d{4})\s*$", header)
    if not m:
        return None
    mes = MESES.get(m.group(1).strip().lower())
    return (int(m.group(2)), mes) if mes else None


def derivar_nombre(path):
    """'Outlook 6+6 2026.XLSX' → 'Outlook 6+6 Unpivot'."""
    stem = os.path.splitext(os.path.basename(path))[0]
    stem = re.sub(r"\s*\b(19|20)\d{2}\b\s*$", "", stem).strip()
    return (stem + " Unpivot") if stem else "Unpivot"


def main():
    ap = argparse.ArgumentParser(description="Genera la hoja «… Unpivot» + columna Valor USD <año>.")
    ap.add_argument("archivo", help="Excel del ejercicio (se modifica in situ)")
    ap.add_argument("--hoja", default=None, help="Hoja de origen (por defecto: la primera)")
    ap.add_argument("--salida", default=None, help="Nombre de la hoja a generar")
    ap.add_argument("--anio-destino", type=int, default=2027, help="Año de referencia USD (def. 2027)")
    ap.add_argument("--vector", default=VECTOR_DEF, help="Vector Ajuste Data Histórica.xlsx")
    ap.add_argument("--serie", default="CPI", help="Fila del vector a usar (def. CPI)")
    args = ap.parse_args()

    src = args.archivo
    if not os.path.exists(src):
        raise SystemExit("ERROR: no existe %s" % src)
    salida = args.salida or derivar_nombre(src)
    dest_y = args.anio_destino

    cpi = leer_cpi(args.vector, args.serie)
    print("Vector: %s · serie %s (%d meses)" % (os.path.basename(args.vector), args.serie, len(cpi)))

    # ---- leer la hoja cruda (valores calculados) ----
    wb_r = openpyxl.load_workbook(src, data_only=True, read_only=True)
    hoja = args.hoja or wb_r.sheetnames[0]
    if hoja not in wb_r.sheetnames:
        raise SystemExit("ERROR: la hoja «%s» no existe. Hojas: %s" % (hoja, wb_r.sheetnames))
    ws_r = wb_r[hoja]

    it = ws_r.iter_rows(values_only=True)
    header = list(next(it))
    meses = []          # (idx_col, año, mes)
    for i, h in enumerate(header):
        pm = parse_mes(h)
        if pm:
            meses.append((i, pm[0], pm[1]))
    if len(meses) != 12:
        print("AVISO: encontré %d columnas de mes (esperaba 12): %s"
              % (len(meses), [header[i] for i, _, _ in meses]))
    if not meses:
        raise SystemExit("ERROR: no encontré columnas de mes tipo «Enero-2026» en «%s»." % hoja)
    src_y = meses[0][1]
    n_id = min(i for i, _, _ in meses)          # las columnas de identidad son las previas al 1er mes
    ident = header[:n_id]
    print("Hoja origen «%s»: %d columnas de identidad + %d meses de %d"
          % (hoja, n_id, len(meses), src_y))

    # ---- factor por mes: CPI(destino-MM) / CPI(origen-MM) ----
    factor = {}
    for _, y, m in meses:
        a, b = cpi.get("%d.%02d" % (dest_y, m)), cpi.get("%d.%02d" % (y, m))
        if a is None or b in (None, 0):
            raise SystemExit("ERROR: falta CPI para %d.%02d o %d.%02d en el vector." % (dest_y, m, y, m))
        factor[m] = a / b
    print("Factores %d→%d: %s" % (src_y, dest_y, ", ".join("%02d:%.6f" % (m, factor[m]) for m in sorted(factor))))

    col_valor = "Valor USD %d (mes equivalente)" % dest_y
    out_header = list(ident) + ["Mes-Año", "Valor", col_valor]

    filas = []
    n_src = 0
    monedas = {}
    try:
        i_mon = [j for j, h in enumerate(ident) if str(h or "").strip().lower().startswith("moneda")][0]
    except IndexError:
        i_mon = None
    suma_mes = {m: 0.0 for _, _, m in meses}
    for row in it:
        if row is None or all(c is None for c in row):
            continue
        n_src += 1
        base = list(row[:n_id])
        if i_mon is not None:
            monedas[base[i_mon]] = monedas.get(base[i_mon], 0) + 1
        for idx, y, m in meses:
            v = row[idx]
            v = 0 if v is None else v
            if not isinstance(v, (int, float)):
                continue
            suma_mes[m] += v
            filas.append(base + ["%d-%02d" % (y, m), v, v * factor[m]])
    wb_r.close()
    print("Filas origen: %d → filas unpivot: %d" % (n_src, len(filas)))
    if monedas:
        print("Monedas:", monedas)
        if any(str(k).upper() != "USD" for k in monedas):
            print("  AVISO: hay filas no-USD; el factor CPI se aplica igual (mismo criterio del archivo de referencia).")

    # ---- escribir la hoja en el MISMO libro (preservando lo existente) ----
    wb_w = openpyxl.load_workbook(src)          # sin data_only → no destruye fórmulas si las hubiera
    if salida in wb_w.sheetnames:
        print("Reemplazo la hoja existente «%s»" % salida)
        del wb_w[salida]
    ws_w = wb_w.create_sheet(title=salida)
    ws_w.append(out_header)
    for f in filas:
        ws_w.append(f)
    ws_w.freeze_panes = "A2"
    wb_w.save(src)
    wb_w.close()

    print("\nOK -> «%s» escrita en %s" % (salida, os.path.basename(src)))
    try:
        print("Control (suma por mes, origen vs unpivot):")
        tot = 0.0
        for _, y, m in meses:
            tot += suma_mes[m]
            print("   %d-%02d  %18s   x%.6f -> %18s"
                  % (y, m, format(suma_mes[m], ",.2f"), factor[m], format(suma_mes[m] * factor[m], ",.2f")))
        print("   %-7s %18s" % ("TOTAL", format(tot, ",.2f")))
    except Exception as e:
        print("   (no pude imprimir el control: %s)" % e)


if __name__ == "__main__":
    main()
