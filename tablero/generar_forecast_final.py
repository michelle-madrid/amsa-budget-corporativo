# -*- coding: utf-8 -*-
"""Genera el Forecast 5+7 2026 FINAL = Forecast 5+7 original + todas las reversas/ajustes.

Junta, en UN solo archivo con el MISMO layout Unpivot que el Forecast original (17 columnas,
hoja «Forecast 5+7 Unpivot»), las filas del Forecast 5+7 original más las filas de cada ajuste
de uploads/ajustes_forecast/ (las mismas que el dashboard suma en la Estructura CECOS
«Nueva con ajustes»). Sumar la columna «Valor» del resultado = Forecast 5+7 con ajustes aplicados.

Salida: salida/ajustes_forecast/Forecast 5+7 2026 (con ajustes).xlsx

Uso:  python generar_forecast_final.py
"""
import os
import openpyxl

import construir as C   # reutiliza FCST_FILE, FCST_SHEET, REV_DIR, _reversas_files, SALIDA_DIR

OUT_DIR = os.path.join(C.SALIDA_DIR, "ajustes_forecast")
OUT_FILE = os.path.join(OUT_DIR, "Forecast 5+7 2026 (con ajustes).xlsx")
OUT_SHEET = C.FCST_SHEET          # «Forecast 5+7 Unpivot» → drop-in del original
NCOLS = 17                        # layout Unpivot (col 0..16); Valor=15, Valor USD 2027=16


def _leer_unpivot(path, sheet):
    """Devuelve (header, [filas]) de una hoja Unpivot, normalizando a NCOLS columnas por posición."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet]
    it = ws.iter_rows(values_only=True)
    header = list(next(it))[:NCOLS]
    filas = []
    for r in it:
        r = list(r)[:NCOLS]
        if len(r) < NCOLS:
            r += [None] * (NCOLS - len(r))
        if not r[C.FC_CECO]:                 # ignora filas sin Centro de Costo (vacías/total)
            continue
        filas.append(r)
    wb.close()
    return header, filas


def _suma(filas, col):
    return sum(v for v in (f[col] for f in filas) if isinstance(v, (int, float)))


def main():
    print("=== Generar Forecast 5+7 FINAL (con ajustes) ===")
    # 1) Forecast 5+7 original
    header, base = _leer_unpivot(C.FCST_FILE, C.FCST_SHEET)
    print(f"Forecast original: {len(base)} filas · Valor={_suma(base, C.FC_VALN):,.0f} · "
          f"USD2027={_suma(base, C.FC_VALA):,.0f}")

    # 2) Cada ajuste/reversa de uploads/ajustes_forecast/ (mismo layout Unpivot)
    ajustes = []
    for path, sheet, label in C._reversas_files():
        _h, filas = _leer_unpivot(path, sheet)
        print(f"  + Ajuste «{label}»: {len(filas)} filas · Valor={_suma(filas, C.FC_VALN):,.0f}")
        ajustes.extend(filas)

    final = base + ajustes
    print(f"FINAL: {len(final)} filas · Valor={_suma(final, C.FC_VALN):,.0f} · "
          f"USD2027={_suma(final, C.FC_VALA):,.0f}")

    # 3) Escribir en el mismo formato Unpivot (header original + todas las filas)
    os.makedirs(OUT_DIR, exist_ok=True)
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet(OUT_SHEET)
    ws.append(header)
    for f in final:
        ws.append(f)
    wb.save(OUT_FILE)
    print("OK -> " + OUT_FILE)


if __name__ == "__main__":
    main()
