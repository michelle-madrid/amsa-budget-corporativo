# -*- coding: utf-8 -*-
"""
revisar_ajustes.py — Audita los ajustes (reversas) al Forecast 5+7.
====================================================================
Lee la carpeta uploads/ajustes_forecast/ EXACTAMENTE como lo hace el build
(construir._reversas_files: cada .xlsx con una hoja «… Unpivot» se incorpora;
el que no la tenga se OMITE) y muestra:

  · Qué archivos están cargados vs. omitidos (y por qué).
  · El acumulado NETO por CECO (Valor USD 2027).
  · El neto total (las redistribuciones suelen dar ~0).

Uso:
  python revisar_ajustes.py                 → resumen: archivos + top CECOs + neto total
  python revisar_ajustes.py 1000AC7010      → detalle de un CECO (por archivo y por CLACO)
  python revisar_ajustes.py --todos         → acumulado de TODOS los CECOs (ordenado)
  python revisar_ajustes.py --excel [ruta]  → escribe un Excel con TODO consolidado
                                              (def.: salida/ajustes_forecast/Ajustes Forecast Consolidado.xlsx)
"""
import os
import sys
from collections import defaultdict

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import construir as B  # reutiliza REV_DIR, FC_* y el mismo auto-descubrimiento

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

MM = 1e6


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def _txt(v):
    return "" if v is None else str(v).strip()


def leer_todo():
    """Devuelve (por_ceco, por_ceco_archivo, por_ceco_claco, archivos, detalle).
    por_ceco[ceco] = suma Valor USD 2027 · por_ceco_archivo[ceco][archivo] · por_ceco_claco[ceco][claco].
    detalle = lista de filas {archivo, ceco, claco, cg, act, vn, va} (agregadas por ceco+claco+archivo)."""
    por_ceco = defaultdict(float)
    por_ceco_archivo = defaultdict(lambda: defaultdict(float))
    por_ceco_claco = defaultdict(lambda: defaultdict(float))
    archivos = []
    # detalle consolidado por (archivo, ceco, claco): suma valores y toma el 1er Concepto/Actividad visto.
    det = {}
    for path, hoja, etiqueta in B._reversas_files():
        fn = os.path.basename(path)
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        it = wb[hoja].iter_rows(values_only=True)
        next(it)  # encabezado
        tot = 0.0
        nfilas = 0
        for r in it:
            if len(r) <= B.FC_VALA:
                continue
            ceco = _txt(r[B.FC_CECO])
            if not ceco:
                continue
            va = _num(r[B.FC_VALA])
            vn = _num(r[B.FC_VALN])
            claco = _txt(r[B.FC_CLACO]) or "(sin claco)"
            por_ceco[ceco] += va
            por_ceco_archivo[ceco][fn] += va
            por_ceco_claco[ceco][claco] += va
            k = (fn, ceco, claco)
            d = det.get(k)
            if d is None:
                det[k] = d = {"archivo": fn, "ceco": ceco, "claco": claco,
                              "cg": _txt(r[B.FC_CG]), "act": _txt(r[B.FC_ACT]), "vn": 0.0, "va": 0.0}
            d["vn"] += vn
            d["va"] += va
            tot += va
            nfilas += 1
        wb.close()
        archivos.append((fn, hoja, nfilas, tot))
    detalle = sorted(det.values(), key=lambda d: (d["ceco"], d["archivo"], d["claco"]))
    return por_ceco, por_ceco_archivo, por_ceco_claco, archivos, detalle


def exportar_excel(ruta, por_ceco, por_ceco_archivo, archivos, detalle):
    """Escribe un .xlsx con 3 hojas: Detalle · Resumen por CECO · Resumen por Archivo."""
    wb = openpyxl.Workbook()
    HDR = Font(bold=True, color="FFFFFF")
    FILL = PatternFill("solid", fgColor="2A8A96")
    NUMFMT = '#,##0'
    def _hoja(ws, headers, filas):
        ws.append(headers)
        for c in ws[1]:
            c.font = HDR; c.fill = FILL; c.alignment = Alignment(horizontal="center")
        for f in filas:
            ws.append(f)
        # formato numérico a las columnas marcadas con None-tipo número: aplicamos por heurística abajo
        return ws

    # Hoja 1 — Detalle (una fila por archivo·CECO·CLACO)
    ws = wb.active; ws.title = "Detalle"
    _hoja(ws, ["Archivo", "CECO", "CLACO", "Concepto Gasto", "Actividad", "Valor Normal USD", "Valor USD 2027"],
          [[d["archivo"], d["ceco"], d["claco"], d["cg"], d["act"], round(d["vn"], 2), round(d["va"], 2)] for d in detalle])
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=7):
        for c in row:
            c.number_format = NUMFMT

    # Hoja 2 — Resumen por CECO (neto)
    ws2 = wb.create_sheet("Resumen por CECO")
    filas = [[ceco, round(v, 2)] for ceco, v in sorted(por_ceco.items(), key=lambda x: -abs(x[1]))]
    _hoja(ws2, ["CECO", "Ajuste neto USD 2027"], filas)
    for row in ws2.iter_rows(min_row=2, min_col=2, max_col=2):
        for c in row:
            c.number_format = NUMFMT
    ws2.append([])
    ws2.append(["NETO TOTAL", round(sum(por_ceco.values()), 2)])
    ws2[ws2.max_row][0].font = Font(bold=True)
    ws2[ws2.max_row][1].font = Font(bold=True); ws2[ws2.max_row][1].number_format = NUMFMT

    # Hoja 3 — Resumen por Archivo
    ws3 = wb.create_sheet("Resumen por Archivo")
    _hoja(ws3, ["Archivo", "Hoja Unpivot", "Filas", "Ajuste neto USD 2027"],
          [[fn, hoja, nfilas, round(tot, 2)] for fn, hoja, nfilas, tot in archivos])
    for row in ws3.iter_rows(min_row=2, min_col=4, max_col=4):
        for c in row:
            c.number_format = NUMFMT

    # Anchos cómodos
    for ws_ in (ws, ws2, ws3):
        for col in ws_.columns:
            letra = col[0].column_letter
            largo = max((len(str(c.value)) for c in col if c.value is not None), default=10)
            ws_.column_dimensions[letra].width = min(46, max(12, largo + 2))

    os.makedirs(os.path.dirname(ruta), exist_ok=True)
    wb.save(ruta)
    return ruta


def main():
    arg = sys.argv[1] if len(sys.argv) > 1 else None
    por_ceco, por_ceco_archivo, por_ceco_claco, archivos, detalle = leer_todo()

    if arg == "--excel":
        ruta = sys.argv[2] if len(sys.argv) > 2 else os.path.join(
            os.path.dirname(B.REV_DIR), "..", "..", "salida", "ajustes_forecast",
            "Ajustes Forecast Consolidado.xlsx")
        ruta = os.path.abspath(ruta)
        exportar_excel(ruta, por_ceco, por_ceco_archivo, archivos, detalle)
        print("✔ Excel consolidado escrito en:")
        print("  " + ruta)
        print(f"  {len(detalle)} filas de detalle · {len(por_ceco)} CECOs · {len(archivos)} archivos.")
        print(f"  Neto total: {sum(por_ceco.values())/MM:+.3f} MM")
        return

    print("=" * 72)
    print("AJUSTES AL FORECAST 5+7  ·  carpeta:", B.REV_DIR)
    print("=" * 72)
    print("Archivos incorporados (con hoja «… Unpivot»):")
    for fn, hoja, nfilas, tot in archivos:
        print(f"  ✔ {fn[:52]:52s} [{hoja[:22]:22s}] {nfilas:5d} filas · {tot/MM:+8.3f} MM")
    if not archivos:
        print("  (ninguno) — deja un .xlsx con hoja «… Unpivot» en la carpeta.")
    # Archivos presentes pero omitidos (sin Unpivot): los detecta _reversas_files con AVISO;
    # acá los listamos comparando el contenido de la carpeta contra lo incorporado.
    cargados = {fn for fn, *_ in archivos}
    todos = [f for f in sorted(os.listdir(B.REV_DIR))
             if f.lower().endswith((".xlsx", ".xlsm")) and not f.startswith("~$")]
    omitidos = [f for f in todos if f not in cargados]
    if omitidos:
        print("\nArchivos OMITIDOS (sin hoja «… Unpivot» → corré unpivot_ejercicio.py):")
        for f in omitidos:
            print("  ✗", f)

    neto = sum(por_ceco.values())
    print(f"\nNeto TOTAL de los ajustes: {neto/MM:+.3f} MM   (≈0 = redistribución pura entre CECOs)")
    print(f"CECOs afectados: {len(por_ceco)}")

    if arg and arg != "--todos":
        ceco = arg
        print("\n" + "=" * 72)
        print(f"DETALLE CECO {ceco}")
        print("=" * 72)
        if ceco not in por_ceco:
            print("  Este CECO no aparece en ningún ajuste.")
            return
        print(f"  Acumulado neto: {por_ceco[ceco]/MM:+.3f} MM")
        print("  Aporte por archivo:")
        for fn, v in sorted(por_ceco_archivo[ceco].items(), key=lambda x: -abs(x[1])):
            print(f"    {v/MM:+8.3f} MM  {fn}")
        print("  Por CLACO:")
        for cl, v in sorted(por_ceco_claco[ceco].items(), key=lambda x: -abs(x[1])):
            if abs(v) >= 0.5:  # oculta ruido sub-USD
                print(f"    {v/MM:+8.3f} MM  CLACO {cl}")
    else:
        top = sorted(por_ceco.items(), key=lambda x: -abs(x[1]))
        top = top if arg == "--todos" else top[:20]
        print("\n" + ("Acumulado por CECO (todos):" if arg == "--todos" else "Top 20 CECOs por |acumulado|:"))
        for ceco, v in top:
            print(f"  {v/MM:+8.3f} MM  {ceco}")
        print("\n(Para el detalle de un CECO:  python revisar_ajustes.py <CECO>)")


if __name__ == "__main__":
    main()
