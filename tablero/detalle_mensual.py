# -*- coding: utf-8 -*-
"""
detalle_mensual.py — Detalle línea a línea CON MES (para el Excel de la BBDD).
=============================================================================
POR QUÉ EXISTE: bbdd/fact_detalle_real.parquet y fact_detalle_ppto.parquet agregan por
BUCKET (2022…2025, 2026 YTD) y por MEDIDA — el mes se pierde ahí, porque el dashboard no
lo necesita a ese nivel. Las FUENTES sí traen el mes ('Mes gasto', 'Mes-Año', y las 12
columnas mensuales del Ppto 2027), así que este módulo las relee y arma el mismo detalle
pero con Año/Mes.

Reusa la lógica de construir.py (mapeo CECO/CLACO→Ítem, remaps, alcance, YTD) importándolo:
las cifras cuadran línea a línea con el dashboard.

Salida (cacheada en bbdd/, se regenera sola si una fuente es más nueva):
  · fact_detalle_real_mes.parquet  → ceco,item,claco,contra,texto_pedido,denominacion,documento,
                                     bucket,anio,mes,valor_n,valor_a,st
  · fact_detalle_ppto_mes.parquet  → ceco,item,claco,concepto_gasto,actividad,medida,ajuste,
                                     anio,mes,valor_n,valor_a

El CLACO del Real no viene en la fuente: se deduce del nombre de la clase de costo con el mismo
diccionario que usa el dashboard (por eso puede quedar vacío en las pocas líneas que no cruzan).

Incluye además las REVERSAS (uploads/ajustes_forecast/) como una medida propia —
'reversas_fcst', con el nombre del archivo en la columna 'ajuste'. En el dashboard esas
reversas solo se suman con la Estructura CECOS «Nueva con ajustes»; acá van aparte para
no pisar el Forecast 5+7 original (Forecast con ajustes = 'forecast_2026' + 'reversas_fcst').

Uso:  python detalle_mensual.py            → arma el caché si hace falta
      python detalle_mensual.py --forzar   → lo rearma siempre
"""
import os
import sys
import openpyxl
import pandas as pd

import construir as C

HERE = os.path.dirname(os.path.abspath(__file__))
BBDD_DIR = C.BBDD_DIR
OUT_REAL = os.path.join(BBDD_DIR, "fact_detalle_real_mes.parquet")
OUT_PPTO = os.path.join(BBDD_DIR, "fact_detalle_ppto_mes.parquet")

# El código también cuenta como fuente: si cambia un mapeo o el corte del YTD, el caché
# tiene que rearmarse aunque los Excel sean los mismos.
_CODIGO = [os.path.join(HERE, "construir.py"), os.path.abspath(__file__), C.CLACOS_FILE]


def _fuentes_real():
    return [C.REAL_FILE, C.DIST_REAL_FILE] + _CODIGO


def _fuentes_ppto():
    """Las reversas se descubren solas (igual que en construir.py), así que se incluye también
    la CARPETA: dejar o sacar un ajuste ahí invalida el caché aunque el .xlsx sea viejo."""
    return ([C.BUD_FILE, C.FCST_FILE, C.OUT_FILE, C.PPTO27_FILE, C.REV_DIR] + _CODIGO
            + [p for p, _h, _l in C._reversas_files()])

MED = {0: "ppto_2025", 1: "ppto_2026ytd", 2: "ppto_2026fy", 3: "forecast_2026",
       4: "ppto_2027", 5: "out66", 6: "reversas_fcst"}
# Ppto 2027: las 12 columnas mensuales (Enero-2027 … Diciembre-2027) van justo antes del Total.
P27_MESES_0 = 14


def _vigente(cache, fuentes):
    """True si ese parquet existe y es más nuevo que TODAS sus fuentes."""
    if not os.path.isfile(cache):
        return False
    t = os.path.getmtime(cache)
    return all(os.path.getmtime(f) <= t for f in fuentes if os.path.exists(f))


# ── Reales ────────────────────────────────────────────────────────────────────────────────
def _real(name2ag, st_names, name2claco, code2item):
    """Mismo detalle que construir.leer_detalle pero agregando por (…, año, mes).
    Mantiene el criterio de alcance temporal: 2022–2025 completos y 2026 solo hasta el
    mes de corte del YTD (construir.YTD_HASTA_MES), igual que el dashboard."""
    agg = {}

    def proc(path, cols):
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        it = wb[C.REAL_SHEET].iter_rows(values_only=True); next(it)
        for r in it:
            ceco = C._txt(r[cols["ceco"]])
            if not ceco:
                continue
            anio = r[cols["anio"]]
            mes = C._txt(r[cols["mes"]]).split(".")[-1].zfill(2)
            if anio in C.BUCKET_IDX:
                bucket = str(anio)
            elif anio == 2026:
                if mes not in C.YTD_2026:      # 2026 entra solo hasta el corte del YTD
                    continue
                bucket = "2026ytd"
            else:
                continue
            nm = C._txt(r[cols["desc"]]).lower()
            item = name2ag.get(nm) or "(sin ítem)"
            # El Real no trae el CLACO como tal: se deduce del nombre de la clase de costo
            # ('Descrip.clases coste') con el mismo diccionario que usa construir.py.
            claco = name2claco.get(nm)
            if claco:
                rm = C._remap_claco(ceco, claco)              # remap fijo de CLACO por CECO
                if rm != claco:                               # solo si el remap aplica, el ítem lo sigue
                    item = code2item.get(rm, item)
                claco = rm
            k = (ceco, item, claco or "", C._txt(r[cols["contra"]]) or "(sin contrapartida)",
                 C._txt(r[cols["texto"]]) or "(sin texto de pedido)",
                 C._txt(r[cols["denom"]]) or "(sin denominación)",
                 C._txt(r[cols["doc"]]) or "—",
                 bucket, int(anio), mes, 1 if nm in st_names else 0)
            vn, va = C._num(r[cols["valn"]]), C._num(r[cols["vala"]])
            cell = agg.get(k)
            if cell:
                cell[0] += vn; cell[1] += va
            else:
                agg[k] = [vn, va]
        wb.close()

    C.log("  Reales corporativos…")
    proc(C.REAL_FILE, C.DET_COLS_CORP)
    C.log("  Reales distribuibles…")
    proc(C.DIST_REAL_FILE, C.DET_COLS_DIST)
    cols = ["ceco", "item", "claco", "contra", "texto_pedido", "denominacion", "documento",
            "bucket", "anio", "mes", "st"]
    return pd.DataFrame([dict(zip(cols, k), valor_n=C._round(v[0]), valor_a=C._round(v[1]))
                         for k, v in agg.items()],
                        columns=cols[:-1] + ["valor_n", "valor_a", "st"])


# ── Ppto / Forecast ───────────────────────────────────────────────────────────────────────
def _ppto(code2ag):
    """Mismo detalle que construir.leer_detalle_pf pero con Año/Mes.
    Ppto 2026 sigue partido en YTD y FY (dos medidas, como hoy): las filas YTD son un
    subconjunto de las FY — NO se suman entre sí."""
    agg = {}

    def add(item, ceco, claco, cg, act, meas, anio, mes, n, a, ajuste=""):
        if n == 0 and a == 0:
            return
        k = (ceco, item, claco or "", cg, act, MED[meas], ajuste, anio, mes)
        c = agg.get(k)
        if c:
            c[0] += n; c[1] += a
        else:
            agg[k] = [n, a]

    C.log("  Planes (Ppto 2025/2026)…")
    wb = openpyxl.load_workbook(C.BUD_FILE, data_only=True, read_only=True)
    it = wb[C.BUD_SHEET].iter_rows(values_only=True); next(it)
    for r in it:
        ceco = C._txt(r[C.B_CECO])
        if not ceco:
            continue
        try:
            anio = int(C._txt(r[C.B_ANIO]))
        except ValueError:
            continue
        if anio not in (2025, 2026):
            continue
        claco = C._remap_claco(ceco, C._txt(r[C.B_CLACO]))
        item = code2ag.get(claco) or "(sin ítem)"
        cg = C._txt(r[C.B_CG]) or "(sin concepto)"; act = C._txt(r[C.B_ACT]) or "(sin actividad)"
        mes = C._txt(r[C.B_MESANIO]).split("-")[-1].zfill(2)
        vn, va = C._num(r[C.B_VALN]), C._num(r[C.B_VALA])
        if anio == 2025:
            add(item, ceco, claco, cg, act, 0, 2025, mes, vn, va)
        else:
            add(item, ceco, claco, cg, act, 2, 2026, mes, vn, va)             # FY
            if mes in C.YTD_2026:
                add(item, ceco, claco, cg, act, 1, 2026, mes, vn, va)         # YTD
    wb.close()

    # Forecast 5+7 (medida 3) y Outlook 6+6 (medida 5): misma hoja Unpivot, col 14 = Mes-Año.
    for path, sheet, meas, lbl in ((C.FCST_FILE, C.FCST_SHEET, 3, "Forecast 5+7"),
                                   (C.OUT_FILE, C.OUT_SHEET, 5, "Outlook 6+6")):
        if not os.path.isfile(path):
            continue
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        if sheet in wb.sheetnames:
            C.log(f"  {lbl}…")
            it = wb[sheet].iter_rows(values_only=True); next(it)
            for r in it:
                ceco = C._txt(r[C.FC_CECO])
                if not ceco:
                    continue
                claco = C._remap_claco(ceco, C._txt(r[C.FC_CLACO]))
                item = code2ag.get(claco) or "(sin ítem)"
                cg = C._txt(r[C.FC_CG]) or "(sin concepto)"
                act = C._txt(r[C.FC_ACT]) or "(sin actividad)"
                add(item, ceco, claco, cg, act, meas, 2026, C._txt(r[14]).split("-")[-1].zfill(2),
                    C._num(r[C.FC_VALN]), C._num(r[C.FC_VALA]))
        wb.close()

    # Reversas / ajustes al Forecast 5+7 (medida 6) — auto-descubiertas en uploads/ajustes_forecast/.
    # Mismo layout Unpivot que el Forecast (reusan los índices FC_*). Van en su PROPIA medida, con el
    # nombre del archivo en 'ajuste': el Forecast con ajustes = forecast_2026 + reversas_fcst.
    for path, sheet, lbl in C._reversas_files():
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        C.log(f"  Reversa: {lbl}…")
        it = wb[sheet].iter_rows(values_only=True); next(it)
        for r in it:
            ceco = C._txt(r[C.FC_CECO])
            if not ceco:
                continue
            claco = C._remap_claco(ceco, C._txt(r[C.FC_CLACO]))
            item = code2ag.get(claco) or "(sin ítem)"
            cg = C._txt(r[C.FC_CG]) or "(sin concepto)"
            act = C._txt(r[C.FC_ACT]) or "(sin actividad)"
            add(item, ceco, claco, cg, act, 6, 2026, C._txt(r[14]).split("-")[-1].zfill(2),
                C._num(r[C.FC_VALN]), C._num(r[C.FC_VALA]), ajuste=lbl)
        wb.close()

    # Ppto 2027 (medida 4): hoja ancha → se despliegan las 12 columnas mensuales.
    # Se controla que la suma de los 12 meses cuadre con la columna Total-2027 (la que usa
    # el dashboard); si no cuadrara, se avisa y se cae al Total como fila única "(anual)".
    if os.path.isfile(C.PPTO27_FILE):
        C.log("  Ppto 2027…")
        wb = openpyxl.load_workbook(C.PPTO27_FILE, data_only=True, read_only=True)
        it = wb[C.PPTO27_SHEET].iter_rows(values_only=True); next(it)
        s_mes = s_tot = 0.0
        pend = []
        for r in it:
            ceco = C._txt(r[C.P27_CECO])
            if not ceco:
                continue
            claco = C._remap_claco(ceco, C._txt(r[C.P27_CLACO]))
            item = code2ag.get(claco) or "(sin ítem)"
            cg = C._txt(r[C.P27_CG]) or "(sin concepto)"
            act = C._txt(r[C.P27_ACT]) or "(sin actividad)"
            meses = [C._num(r[P27_MESES_0 + i]) for i in range(12)]
            s_mes += sum(meses); s_tot += C._num(r[C.P27_TOTAL])
            pend.append((item, ceco, claco, cg, act, meses, C._num(r[C.P27_TOTAL])))
        wb.close()
        cuadra = abs(s_mes - s_tot) <= max(1.0, abs(s_tot) * 1e-6)
        if not cuadra:
            C.log(f"  OJO Ppto 2027: los 12 meses suman {s_mes:,.0f} y el Total {s_tot:,.0f} "
                  f"→ se usa el Total en una fila «(anual)».")
        for item, ceco, claco, cg, act, meses, tot in pend:
            if cuadra:
                for i, v in enumerate(meses):
                    add(item, ceco, claco, cg, act, 4, 2027, f"{i + 1:02d}", v, v)
            else:
                add(item, ceco, claco, cg, act, 4, 2027, "(anual)", tot, tot)

    cols = ["ceco", "item", "claco", "concepto_gasto", "actividad", "medida", "ajuste", "anio", "mes"]
    return pd.DataFrame([dict(zip(cols, k), valor_n=C._round(v[0]), valor_a=C._round(v[1]))
                         for k, v in agg.items()],
                        columns=cols + ["valor_n", "valor_a"])


def construir(forzar=False):
    """Devuelve (df_real, df_ppto) con Año/Mes. Cada mitad usa su caché si sigue vigente:
    tocar el Ppto/Forecast NO obliga a releer los Reales (que son los archivos pesados)."""
    hay_real = not forzar and _vigente(OUT_REAL, _fuentes_real())
    hay_ppto = not forzar and _vigente(OUT_PPTO, _fuentes_ppto())
    if hay_real and hay_ppto:
        return pd.read_parquet(OUT_REAL), pd.read_parquet(OUT_PPTO)

    C.log("Armando detalle MENSUAL desde las fuentes" +
          (" (Reales: tarda unos minutos)" if not hay_real else " (solo Ppto/Forecast)") + "…")
    code2item, name2item, *_rest = C.leer_clacos()
    st_names, name2claco = _rest[5], _rest[6]
    os.makedirs(BBDD_DIR, exist_ok=True)

    if hay_real:
        dr = pd.read_parquet(OUT_REAL)
    else:
        dr = _real(name2item, st_names, name2claco, code2item)
        dr.to_parquet(OUT_REAL, index=False)
        C.log(f"  bbdd/fact_detalle_real_mes.parquet · {len(dr):,} filas")
    if hay_ppto:
        dp = pd.read_parquet(OUT_PPTO)
    else:
        dp = _ppto(code2item)
        dp.to_parquet(OUT_PPTO, index=False)
        C.log(f"  bbdd/fact_detalle_ppto_mes.parquet · {len(dp):,} filas")
    return dr, dp


if __name__ == "__main__":
    construir(forzar="--forzar" in sys.argv)
